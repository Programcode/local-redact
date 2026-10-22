"""Robustness: hundreds of mixed-format synthetic files + fail-closed skips.

Uses StubEngine (no model load). Requires text_fmt/csv handlers at minimum;
skips individual format checks if their modules are missing.
"""

from __future__ import annotations

import csv
from pathlib import Path

import pytest

from local_redact.batch import batch_redact
from tests.conftest import StubEngine


def _has(mod: str) -> bool:
    try:
        __import__(f"local_redact.formats.{mod}")
        return True
    except Exception:
        return False


@pytest.mark.slow
def test_hundreds_of_mixed_files(tmp_path, stub_engine):
    if not _has("text_fmt"):
        pytest.skip("text_fmt required")

    eng = stub_engine([("SECRETNAME", "PERSON_00000")])
    src = tmp_path / "corpus"
    dst = tmp_path / "out"
    src.mkdir()

    # 200 plain text files
    for i in range(200):
        (src / f"doc_{i:03d}.txt").write_text(
            f"Report {i}: SECRETNAME present.\n", encoding="utf-8"
        )

    # 50 markdown
    for i in range(50):
        (src / f"note_{i:02d}.md").write_text(
            f"# Note {i}\n\nSECRETNAME details.\n", encoding="utf-8"
        )

    # 30 csv
    for i in range(30):
        with (src / f"sheet_{i:02d}.csv").open("w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["id", "owner"])
            w.writerow([i, "SECRETNAME"])

    # unsupported binary junk — must be skipped, not copied
    for i in range(20):
        (src / f"junk_{i:02d}.bin").write_bytes(b"\x00SECRETNAME\xff")

    # optional richer formats if handlers exist
    if _has("docx_fmt"):
        from docx import Document

        for i in range(10):
            d = Document()
            d.add_paragraph(f"Owner is SECRETNAME file {i}")
            d.save(src / f"word_{i:02d}.docx")

    if _has("xlsx_fmt"):
        from openpyxl import Workbook

        for i in range(10):
            wb = Workbook()
            ws = wb.active
            ws["A1"] = "SECRETNAME"
            ws["B1"] = f"=A1&{i}"
            wb.save(src / f"xl_{i:02d}.xlsx")

    if _has("pptx_fmt"):
        from pptx import Presentation

        for i in range(5):
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            box = slide.shapes.add_textbox(0, 0, 3000000, 500000)
            box.text_frame.text = f"SECRETNAME slide {i}"
            prs.save(src / f"deck_{i:02d}.pptx")

    if _has("pdf_fmt"):
        import fitz

        for i in range(5):
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((72, 72), f"SECRETNAME pdf {i}")
            doc.save(src / f"pdf_{i:02d}.pdf")
            doc.close()

    report = batch_redact(src, dst, eng)
    summary = report.counts
    assert summary.get("redacted", 0) + summary.get("partial", 0) >= 280
    assert summary.get("skipped", 0) >= 20
    assert summary.get("failed", 0) in (0, None) or summary.get("failed", 0) == 0

    # nothing unredacted lies in output as bare SECRETNAME for redacted files
    for path in dst.rglob("*"):
        if path.is_file() and path.suffix.lower() in {".txt", ".md", ".csv"}:
            body = path.read_text(encoding="utf-8")
            assert "SECRETNAME" not in body, path

    # junk never copied
    assert not list(dst.glob("junk_*.bin"))
    assert (dst / "redaction_report.json").exists()
