"""Tests for the XLSX handler.

Critical invariant under test: formulas are NEVER rewritten (a redacted
formula breaks the sheet) — only literal string cell values are redacted.
Numbers/dates pass through untouched.
"""

from __future__ import annotations

from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

try:
    from local_redact.formats.xlsx_fmt import XlsxHandler
except ImportError:
    # formats/__init__.py eagerly imports sibling handlers (text_fmt,
    # docx_fmt, ...) owned by other workstreams. While they have not landed,
    # load xlsx_fmt through a bare package shim so this suite runs alone.
    # Once all siblings exist the normal import above succeeds and this
    # branch never executes.
    import sys
    import types
    from pathlib import Path

    import local_redact

    _shim = types.ModuleType("local_redact.formats")
    _shim.__path__ = [
        str(Path(local_redact.__file__).resolve().parent / "formats")
    ]
    sys.modules["local_redact.formats"] = _shim
    from local_redact.formats.xlsx_fmt import XlsxHandler


def _make_workbook(path, sheets: dict[str, dict[str, object]]) -> None:
    """Write an xlsx at `path` with {sheet_name: {cell_ref: value}}."""
    wb = Workbook()
    # Repurpose the default sheet for the first entry, add the rest.
    first = True
    for name, cells in sheets.items():
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(name)
        for ref, value in cells.items():
            ws[ref] = value
    wb.save(path)
    wb.close()


def _redact(tmp_path, stub_engine, sheets, pairs):
    src = tmp_path / "in.xlsx"
    dst = tmp_path / "out.xlsx"
    _make_workbook(src, sheets)
    engine = stub_engine(pairs)
    report = XlsxHandler().redact(src, dst, engine, "docs/in.xlsx")
    assert dst.exists(), "handler must always write the output file"
    return report, load_workbook(dst)


def test_extensions():
    assert XlsxHandler.extensions == (".xlsx",)


def test_string_cell_redacted(tmp_path, stub_engine):
    report, wb = _redact(
        tmp_path,
        stub_engine,
        {"Sheet1": {"A1": "John Doe", "B1": "no entities here"}},
        [("John Doe", "[PERSON_1]")],
    )
    ws = wb["Sheet1"]
    assert ws["A1"].value == "[PERSON_1]"
    assert ws["B1"].value == "no entities here"
    assert report.status == "redacted"
    assert report.counts == {"PERSON": 1}
    wb.close()


def test_redaction_inside_longer_string(tmp_path, stub_engine):
    report, wb = _redact(
        tmp_path,
        stub_engine,
        {"Sheet1": {"A1": "Contact John Doe or John Doe's office"}},
        [("John Doe", "[PERSON_1]")],
    )
    assert wb["Sheet1"]["A1"].value == "Contact [PERSON_1] or [PERSON_1]'s office"
    assert report.counts == {"PERSON": 2}
    wb.close()


def test_formula_cells_never_rewritten(tmp_path, stub_engine):
    """The CRITICAL spec rule: a redacted formula breaks the sheet."""
    report, wb = _redact(
        tmp_path,
        stub_engine,
        {
            "Sheet1": {
                "A1": 10,
                "B1": 20,
                "C1": "=A1+B1",
                # PII literal embedded in a formula: still untouched.
                "C2": '="John Doe"',
                "A2": "John Doe",
            }
        },
        [("John Doe", "[PERSON_1]")],
    )
    ws = wb["Sheet1"]
    assert ws["C1"].value == "=A1+B1"
    assert ws["C2"].value == '="John Doe"'
    # The plain string cell with the same value IS redacted.
    assert ws["A2"].value == "[PERSON_1]"
    # Formula text was never sent through the engine: only 1 replacement.
    assert report.counts == {"PERSON": 1}
    wb.close()


def test_non_string_cells_untouched(tmp_path, stub_engine):
    moment = datetime(2024, 1, 15, 9, 30)
    report, wb = _redact(
        tmp_path,
        stub_engine,
        {"Sheet1": {"A1": 42, "A2": 3.14, "A3": moment, "A4": True, "A5": None}},
        [("42", "[NUM]"), ("John Doe", "[PERSON_1]")],
    )
    ws = wb["Sheet1"]
    assert ws["A1"].value == 42
    assert ws["A2"].value == 3.14
    assert ws["A3"].value == moment
    assert ws["A4"].value is True
    assert ws["A5"].value is None
    assert report.counts == {}
    wb.close()


def test_all_sheets_processed(tmp_path, stub_engine):
    report, wb = _redact(
        tmp_path,
        stub_engine,
        {
            "First": {"A1": "John Doe"},
            "Second": {"B2": "Jane Smith"},
            "Third": {"C3": "nothing"},
        },
        [("John Doe", "[PERSON_1]"), ("Jane Smith", "[PERSON_2]")],
    )
    assert wb["First"]["A1"].value == "[PERSON_1]"
    assert wb["Second"]["B2"].value == "[PERSON_2]"
    assert wb["Third"]["C3"].value == "nothing"
    assert report.counts == {"PERSON": 2}
    wb.close()


def test_no_matches_still_writes_output(tmp_path, stub_engine):
    report, wb = _redact(
        tmp_path,
        stub_engine,
        {"Sheet1": {"A1": "plain text", "B1": 7}},
        [("John Doe", "[PERSON_1]")],
    )
    assert report.status == "redacted"
    assert report.counts == {}
    assert wb["Sheet1"]["A1"].value == "plain text"
    assert wb["Sheet1"]["B1"].value == 7
    wb.close()


def test_report_rel_path(tmp_path, stub_engine):
    report, wb = _redact(
        tmp_path, stub_engine, {"Sheet1": {"A1": "x"}}, []
    )
    assert report.rel_path == "docs/in.xlsx"
    wb.close()


def test_corrupt_file_raises(tmp_path, stub_engine):
    """Fail-closed contract: handler raises, dispatcher marks failed and
    never copies the file to the output tree."""
    src = tmp_path / "broken.xlsx"
    src.write_bytes(b"this is not a zip archive at all")
    dst = tmp_path / "out.xlsx"
    with pytest.raises(Exception):
        XlsxHandler().redact(src, dst, stub_engine([]), "docs/broken.xlsx")
