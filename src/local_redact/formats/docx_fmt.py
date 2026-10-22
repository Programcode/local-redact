"""DOCX handler (python-docx).

Walks paragraphs and table cells; redacts each paragraph's full text via a
single engine.redact_text call, then puts the result into the first run and
clears subsequent runs so layout/paragraph structure remains but split-run
formatting on a redacted span collapses to the first run's style (acceptable
for Phase 1 — redaction inherently disrupts character-level formatting).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

from .base import FileReport, Handler


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Replace paragraph text while keeping at least one run for style."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _iter_paragraphs(doc: Document):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
    # headers / footers
    for section in doc.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            if part is None:
                continue
            for p in part.paragraphs:
                yield p
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            yield p


class DocxHandler(Handler):
    extensions = (".docx",)

    def redact(self, src: Path, dst: Path, engine, rel_doc: str) -> FileReport:
        doc = Document(str(src))
        counts: dict[str, int] = {}
        for paragraph in _iter_paragraphs(doc):
            original = paragraph.text
            if not original:
                continue
            result = engine.redact_text(original, doc=rel_doc, audit=False)
            if result.text != original:
                _set_paragraph_text(paragraph, result.text)
            for etype, n in result.counts_by_type.items():
                counts[etype] = counts.get(etype, 0) + n
        doc.save(str(dst))
        return FileReport(rel_path=rel_doc, status="redacted", counts=counts)
